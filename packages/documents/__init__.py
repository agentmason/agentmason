from __future__ import annotations

import csv
import json
import logging
from abc import ABC, abstractmethod
from typing import Optional

logger = logging.getLogger(__name__)


class ParseError(Exception):
    """Exception raised during document parsing."""
    pass


class ParseResult:
    """Result of parsing a document."""

    def __init__(
        self,
        text: str,
        metadata: Optional[dict] = None,
        pages: Optional[int] = None,
    ) -> None:
        self.text = text
        self.metadata = metadata or {}
        self.pages = pages
        if pages:
            self.metadata["total_pages"] = pages


class DocumentParser(ABC):
    """Abstract base class for document parsers."""

    @staticmethod
    @abstractmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        """
        Parse a document file and extract text.
        
        Args:
            file_path: Path to the file
            file_name: Original file name
            
        Returns:
            ParseResult containing extracted text and metadata
            
        Raises:
            ParseError: If parsing fails
        """
        pass

    @staticmethod
    @abstractmethod
    def supports_mime_type(mime_type: str) -> bool:
        """Check if parser supports the given MIME type."""
        pass


class TextParser(DocumentParser):
    """Parser for plain text files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            return ParseResult(text=text, metadata={"file_name": file_name})
        except Exception as e:
            raise ParseError(f"Failed to parse text file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in [
            "text/plain",
            "text/markdown",
            "text/x-markdown",
        ]


class MarkdownParser(DocumentParser):
    """Parser for Markdown files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            return ParseResult(text=text, metadata={"file_name": file_name})
        except Exception as e:
            raise ParseError(f"Failed to parse markdown file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in ["text/markdown", "text/x-markdown"]


class CSVParser(DocumentParser):
    """Parser for CSV files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            rows = []
            with open(file_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                if not reader.fieldnames:
                    raise ParseError("CSV file is empty or has no headers")

                for row in reader:
                    # Convert each row to readable text format
                    row_text = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
                    rows.append(row_text)

            text = "\n".join(rows)
            return ParseResult(
                text=text,
                metadata={
                    "file_name": file_name,
                    "format": "csv",
                    "columns": reader.fieldnames,
                }
            )
        except Exception as e:
            raise ParseError(f"Failed to parse CSV file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in [
            "text/csv",
            "application/csv",
            "application/x-csv",
        ]


class JSONParser(DocumentParser):
    """Parser for JSON files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            # Convert JSON to readable text format
            text = json.dumps(data, indent=2)
            return ParseResult(text=text, metadata={"file_name": file_name, "format": "json"})
        except json.JSONDecodeError as e:
            raise ParseError(f"Invalid JSON file: {str(e)}") from e
        except Exception as e:
            raise ParseError(f"Failed to parse JSON file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in [
            "application/json",
            "application/x-json",
        ]


class PDFParser(DocumentParser):
    """Parser for PDF files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            import PyPDF2
        except ImportError:
            raise ParseError("PyPDF2 is required for PDF parsing. Install with: pip install PyPDF2")

        try:
            text_parts = []
            page_count = 0

            with open(file_path, "rb") as f:
                pdf_reader = PyPDF2.PdfReader(f)
                page_count = len(pdf_reader.pages)

                for page_num, page in enumerate(pdf_reader.pages, 1):
                    text = page.extract_text()
                    if text:
                        # Add page marker for later filtering
                        text_parts.append(f"[Page {page_num}]\n{text}")

            if not text_parts:
                raise ParseError("No text could be extracted from PDF")

            text = "\n\n".join(text_parts)
            return ParseResult(
                text=text,
                metadata={"file_name": file_name, "format": "pdf"},
                pages=page_count
            )
        except ParseError:
            raise
        except Exception as e:
            raise ParseError(f"Failed to parse PDF file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in ["application/pdf"]


class DOCXParser(DocumentParser):
    """Parser for DOCX files."""

    @staticmethod
    async def parse(file_path: str, file_name: str) -> ParseResult:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ParseError("python-docx is required for DOCX parsing. Install with: pip install python-docx")

        try:
            doc = DocxDocument(file_path)
            text_parts = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)

            if not text_parts:
                raise ParseError("No text could be extracted from DOCX")

            text = "\n".join(text_parts)
            return ParseResult(text=text, metadata={"file_name": file_name, "format": "docx"})
        except ParseError:
            raise
        except Exception as e:
            raise ParseError(f"Failed to parse DOCX file: {str(e)}") from e

    @staticmethod
    def supports_mime_type(mime_type: str) -> bool:
        return mime_type in [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ]


class ParserRegistry:
    """Registry for document parsers."""

    def __init__(self) -> None:
        self._parsers: list[type[DocumentParser]] = [
            TextParser,
            MarkdownParser,
            CSVParser,
            JSONParser,
            PDFParser,
            DOCXParser,
        ]

    def get_parser(self, mime_type: str) -> type[DocumentParser]:
        """Get appropriate parser for MIME type."""
        for parser_class in self._parsers:
            if parser_class.supports_mime_type(mime_type):
                return parser_class

        raise ParseError(f"No parser available for MIME type: {mime_type}")

    async def parse(self, file_path: str, file_name: str, mime_type: str) -> ParseResult:
        """Parse a document with appropriate parser."""
        parser_class = self.get_parser(mime_type)
        return await parser_class.parse(file_path, file_name)
